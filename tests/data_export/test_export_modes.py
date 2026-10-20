"""TEST_ONLY 导表端到端回归：可重复、保留正式数据、失败原子性、禁止误用模式。"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
TOOL = ROOT / "tools/export_game_data.py"
FIXTURE = ROOT / "tests/fixtures/data_export/test_only_data.xlsx"


def execute(
    path: Path, *flags: str, output_root: Path | None = None,
) -> subprocess.CompletedProcess:
    """运行真实入口，并锁定 UTF-8 控制台输出，避免 Windows 本地代码页乱码。"""
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    command = [sys.executable, str(TOOL), "--input", str(path)]
    if output_root is not None:
        command.extend(("--output-root", str(output_root)))
    command.extend(flags)
    return subprocess.run(
        command,
        cwd=ROOT, capture_output=True, encoding="utf-8", env=env,
        check=False,
    )


def files_snapshot(paths: list[Path]) -> dict[Path, tuple[bytes, int]]:
    """用文件字节和 mtime 判定重复导表是否真正稳定。"""
    return {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in paths}


def main() -> None:
    # 所有导出写入独立临时根目录，同时核对仓库正式来源与 Tier Resource 未变。
    authoritative_paths = sorted((ROOT / "data/source_tables").glob("*.csv"))
    for name in ("01_身份配置.csv", "06_战斗数值.csv", "08_Tier档位.csv"):
        assert ROOT / "data/source_tables" / name in authoritative_paths, f"缺少已确认来源：{name}"
    authoritative_paths.append(ROOT / "data/combat_stage/tier_catalog.tres")
    authoritative = files_snapshot(authoritative_paths)
    with tempfile.TemporaryDirectory() as temp:
        temp_root = Path(temp)
        test_root = temp_root / "test-output"

        success = execute(FIXTURE, "--test-only", output_root=test_root)
        assert success.returncode == 0, (success.stdout, success.stderr)
        test_outputs = sorted((test_root / "data/test_only/source_tables").glob("*.csv"))
        test_outputs += sorted((test_root / "data/test_only/generated").rglob("*.tres"))
        assert len(test_outputs) == 20, len(test_outputs)
        state = files_snapshot(test_outputs)
        success = execute(FIXTURE, "--test-only", output_root=test_root)
        assert success.returncode == 0, (success.stdout, success.stderr)
        assert state == files_snapshot(test_outputs)
        assert authoritative == files_snapshot(authoritative_paths)
        print("PASS TEST_ONLY repeatable output: 15 CSV + 5 Resource; source unchanged")

        # A3 原有 test_word_01，复制到第二条数据制造重复 ID。
        wb = load_workbook(FIXTURE)
        wb["03_普通词库"]["A4"] = "test_word_01"
        duplicate = temp_root / "duplicate.xlsx"
        wb.save(duplicate)
        wb.close()
        result = execute(duplicate, "--test-only", output_root=test_root)
        assert result.returncode == 1 and "稳定 ID 重复" in result.stderr
        assert state == files_snapshot(test_outputs), "错误时覆盖了旧版测试资产"
        print("PASS duplicate ID rejected without replacing existing assets")

        # 词库引用有双方来源，拼错关卡的 pool_id 必须失败并保留旧产物。
        wb = load_workbook(FIXTURE)
        sheet = wb["02_主播关卡"]
        header = next(row for row in sheet.iter_rows() if any(cell.value == "word_pool_id" for cell in row))
        column = next(cell.column for cell in header if cell.value == "word_pool_id")
        sheet.cell(header[0].row + 1, column, "test_missing_pool")
        dangling = temp_root / "dangling_pool.xlsx"
        wb.save(dangling)
        wb.close()
        result = execute(dangling, "--test-only", output_root=test_root)
        assert result.returncode == 1 and "关联 ID" in result.stderr, (result.stdout, result.stderr)
        assert state == files_snapshot(test_outputs), "跨表错误覆盖了旧版测试资产"
        assert authoritative == files_snapshot(authoritative_paths), "失败导出改变了正式来源"
        print("PASS missing cross-table pool rejected without replacing assets")

        wb = load_workbook(FIXTURE)
        wb["00_填写说明"]["A1"] = "普通数据表"
        non_test = temp_root / "non_test.xlsx"
        wb.save(non_test)
        wb.close()
        result = execute(non_test, "--test-only", output_root=test_root)
        assert result.returncode == 2 and "--test-only" in result.stderr
        print("PASS prevents official/unmarked workbook entering TEST_ONLY mode")

        # 普通模式拒绝 TEST_ONLY 来源时，既存 CSV 的字节、mtime 与目录文件集合都保持。
        normal_root = temp_root / "normal-rejection-output"
        source_dir = normal_root / "data/source_tables"
        source_dir.mkdir(parents=True)
        previous = source_dir / "02_主播关卡.csv"
        legacy = source_dir / "legacy.csv"
        previous.write_bytes(b"level_id,level_order\r\nold_level,1\r\n")
        legacy.write_bytes(b"legacy\r\n")
        preserved_mtime = 1_700_000_000_123_456_789
        os.utime(previous, ns=(preserved_mtime, preserved_mtime))
        before_rejection = files_snapshot([previous, legacy])
        rejected = execute(FIXTURE, output_root=normal_root)
        assert rejected.returncode == 2, (rejected.returncode, rejected.stdout, rejected.stderr)
        assert "普通导表模式拒绝" in rejected.stderr and "TEST_ONLY" in rejected.stderr
        assert files_snapshot([previous, legacy]) == before_rejection
        assert sorted(path.name for path in source_dir.glob("*.csv")) == ["02_主播关卡.csv", "legacy.csv"]
        print("PASS ordinary mode rejects TEST_ONLY before writes: exit=2; existing CSV bytes/mtime preserved")

        # 将现有结构夹具转成未标记、字段关联有效的普通输入，验证正式输出路径仍可用。
        wb = load_workbook(FIXTURE)
        for worksheet in wb.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = cell.value.replace("TEST_ONLY", "正式").replace("test_", "formal_")
        wb["00_填写说明"]["A1"] = "正式策划工作簿"
        official_input = temp_root / "unmarked_formal.xlsx"
        wb.save(official_input)
        wb.close()
        official_root = temp_root / "normal-success-output"
        normal_success = execute(official_input, output_root=official_root)
        assert normal_success.returncode == 0, (normal_success.stdout, normal_success.stderr)
        official_csvs = sorted((official_root / "data/source_tables").glob("*.csv"))
        assert len(official_csvs) == 20, len(official_csvs)
        assert "formal_level_01" in (official_root / "data/source_tables/02_主播关卡.csv").read_text(encoding="utf-8")
        assert not (official_root / "data/test_only/source_tables").exists()
        assert authoritative == files_snapshot(authoritative_paths)
        print("PASS validated unmarked workbook uses ordinary source path: exit=0; 20 CSV; repository source unchanged")

        # 真实 V0 规则：T0 无回拉可设为零，停用的旧吞并奖励示例可保留旧主播 ID。
        wb = load_workbook(official_input)
        tier_ws = wb["08_Tier档位"]
        tier_header = {cell.value: cell.column for cell in tier_ws[1]}
        tier_col = tier_header["tier"]
        pullback_col = tier_header["pullback_mult"]
        tier_rows = {}
        for row in range(2, tier_ws.max_row + 1):
            tier_value = tier_ws.cell(row, tier_col).value
            if tier_value is not None:
                tier_rows[str(tier_value).strip()] = row
        assert "0" in tier_rows and "1" in tier_rows
        tier_ws.cell(tier_rows["0"], pullback_col, 0)
        reward_ws = wb["13_吞并奖励"]
        reward_header = [cell.value for cell in reward_ws[1]]
        placeholder_values = {
            "reward_id": "formal_disabled_legacy_reward",
            "streamer_id": "former_streamer_removed",
            "reward_type": "word_pool", "content_id": "formal_legacy_pool",
            "inherited_weight": "1", "enabled": "FALSE", "notes": "停用旧记录回归",
        }
        reward_ws.append([placeholder_values.get(name, "") for name in reward_header])
        reward_enabled_col = reward_header.index("enabled") + 1
        reward_row = reward_ws.max_row
        v0_input = temp_root / "official_v0.xlsx"
        wb.save(v0_input)
        wb.close()
        v0_root = temp_root / "official_v0-output"
        v0_result = execute(v0_input, output_root=v0_root)
        assert v0_result.returncode == 0, (v0_result.stdout, v0_result.stderr)
        assert (v0_root / "data/source_tables/08_Tier档位.csv").exists()
        print("PASS T0 zero pullback and disabled legacy reward accepted")

        wb = load_workbook(v0_input)
        wb["13_吞并奖励"].cell(reward_row, reward_enabled_col, "TRUE")
        enabled_bad = temp_root / "enabled_bad_reward.xlsx"
        wb.save(enabled_bad)
        wb.close()
        enabled_result = execute(enabled_bad, output_root=temp_root / "enabled-bad-output")
        assert enabled_result.returncode == 1 and "关联 ID" in enabled_result.stderr, (enabled_result.stdout, enabled_result.stderr)
        print("PASS enabled dangling reward rejected")

        wb = load_workbook(v0_input)
        wb["08_Tier档位"].cell(tier_rows["1"], pullback_col, 0)
        bad_tier = temp_root / "tier1_zero_pullback.xlsx"
        wb.save(bad_tier)
        wb.close()
        tier_result = execute(bad_tier, output_root=temp_root / "tier-bad-output")
        assert tier_result.returncode == 1 and "pullback_mult" in tier_result.stderr, (tier_result.stdout, tier_result.stderr)
        print("PASS Tier1 zero pullback rejected")

        # EN-10 regression: formal CSV import must preserve the approved ending card-back mapping.
        approved_root = temp_root / "approved-csv-output"
        approved_env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
        approved_result = subprocess.run(
            [sys.executable, str(TOOL), "--approved-csv", "--output-root", str(approved_root)],
            cwd=ROOT, capture_output=True, encoding="utf-8", env=approved_env, check=False,
        )
        assert approved_result.returncode == 0, (approved_result.stdout, approved_result.stderr)
        generated_art = (approved_root / "data/ending/ending_main_art_config.tres").read_text(encoding="utf-8")
        committed_art = (ROOT / "data/ending/ending_main_art_config.tres").read_text(encoding="utf-8")
        card_back_ref = 'path="res://assets/ui/identity/cards/identity_card_back.png"'
        assert generated_art == committed_art, "checked-in ending Resource drifted from formal CSV export"
        assert generated_art.count(card_back_ref) == 3, generated_art
        for tendency in ("orthodox", "heretical", "absurd"):
            assert f'{tendency}_main_art = ExtResource("art_{tendency}")' in generated_art, generated_art
        assert "_main_art = null" not in generated_art, generated_art
        print("PASS EN-10 approved CSV re-export keeps all three ending art slots bound to the shared card back")


if __name__ == "__main__":
    main()
